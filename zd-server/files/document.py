#!/usr/bin/env python3
"""
Чтение документа компании роботом Google (только чтение).

  python3 document.py <ссылка или id>            текст документа
  python3 document.py <ссылка> --вкладки         вкладки таблицы
  python3 document.py <ссылка> --лист <вкладка>  данные вкладки

Таблицы читаются через Sheets API по вкладкам: формулы посчитаны, структура цела.
Выгрузка в csv ломает финмодели, поэтому так не делаем.
"""
import os, re, sys, io, json, warnings
warnings.filterwarnings("ignore")
import requests
from google.oauth2 import service_account
from google.auth.transport.requests import Request

SA = os.environ.get("ROBOT_KEY", os.path.expanduser("~/.aplayers/robot.json"))
DRIVE = "https://www.googleapis.com/drive/v3/files"
SHEETS = "https://sheets.googleapis.com/v4/spreadsheets"

def H():
    c = service_account.Credentials.from_service_account_file(SA, scopes=[
        "https://www.googleapis.com/auth/drive.readonly",
        "https://www.googleapis.com/auth/spreadsheets.readonly"])
    c.refresh(Request()); return {"Authorization": "Bearer " + c.token}

def fid(s):
    m = re.search(r"/(?:d|folders)/([A-Za-z0-9_-]{20,})", s or "")
    return m.group(1) if m else s

def main():
    if len(sys.argv) < 2: print(__doc__); return
    h = H(); i = fid(sys.argv[1])
    meta = requests.get(f"{DRIVE}/{i}", headers=h,
        params={"supportsAllDrives": "true", "fields": "id,name,mimeType"}, timeout=30)
    if meta.status_code != 200:
        sys.exit(f"Робот не видит этот документ ({meta.status_code}). "
                 f"Значит, его не расшарили роботу. Так и скажи, не выдумывай содержимое.")
    m = meta.json(); mt = m["mimeType"]
    print(f"# {m['name']}\n")
    if mt.endswith("spreadsheet"):
        if "--вкладки" in sys.argv:
            r = requests.get(f"{SHEETS}/{i}", headers=h,
                params={"fields": "sheets.properties(title,gridProperties)"}, timeout=30)
            for s in r.json().get("sheets", []):
                p = s["properties"]; g = p.get("gridProperties", {})
                print(f"  · {p['title']}  ({g.get('rowCount')}x{g.get('columnCount')})")
            return
        tab = None
        if "--лист" in sys.argv: tab = sys.argv[sys.argv.index("--лист") + 1]
        if not tab:
            r = requests.get(f"{SHEETS}/{i}", headers=h, params={"fields": "sheets.properties.title"}, timeout=30)
            tabs = [s["properties"]["title"] for s in r.json().get("sheets", [])]
            print("Вкладки:", ", ".join(tabs))
            print("\nЧтобы прочитать: document.py <ссылка> --лист <вкладка>")
            return
        import urllib.parse
        a1 = urllib.parse.quote(f"{tab}!A1:BZ400")
        r = requests.get(f"{SHEETS}/{i}/values/{a1}", headers=h,
            params={"valueRenderOption": "UNFORMATTED_VALUE", "dateTimeRenderOption": "FORMATTED_STRING"}, timeout=40)
        for row in r.json().get("values", []):
            print("\t".join("" if c is None else str(c) for c in row))
        return
    if mt.endswith("document"):
        r = requests.get(f"{DRIVE}/{i}/export", headers=h, params={"mimeType": "text/plain"}, timeout=60)
        r.encoding = "utf-8"          # Google отдает utf-8, но без заголовка кодировки
        print(r.text.lstrip("\ufeff")[:60000]); return
    if mt.endswith("presentation"):
        r = requests.get(f"{DRIVE}/{i}/export", headers=h, params={"mimeType": "text/plain"}, timeout=60)
        r.encoding = "utf-8"
        print(r.text.lstrip("\ufeff")[:60000]); return
    if mt.endswith("folder"):
        r = requests.get(DRIVE, headers=h, params={"supportsAllDrives": "true",
            "includeItemsFromAllDrives": "true", "q": f"'{i}' in parents and trashed=false",
            "fields": "files(id,name,mimeType)"}, timeout=30)
        for f in r.json().get("files", []):
            print(f"  {f['name']}\n    https://drive.google.com/file/d/{f['id']}")
        return
    r = requests.get(f"{DRIVE}/{i}", headers=h, params={"supportsAllDrives": "true", "alt": "media"}, timeout=90)
    data = r.content
    if mt.endswith("spreadsheetml.sheet"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            print(f"\n## {ws.title}")
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n > 200: print("... обрезано"); break
                if any(c is not None for c in row):
                    print("\t".join("" if c is None else str(c) for c in row))
    elif mt.endswith("wordprocessingml.document"):
        import docx
        for p in docx.Document(io.BytesIO(data)).paragraphs:
            if p.text.strip(): print(p.text)
    else:
        print(f"(файл {mt}, {len(data)} байт, текстом не читается)")

if __name__ == "__main__":
    main()
